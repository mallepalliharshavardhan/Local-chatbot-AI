import json
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from src.application_tracker import ApplicationTracker
from src.aihawk_job_manager import AIHawkJobManager


def test_tracker_creates_and_updates_application_rows(tmp_path):
    tracker = ApplicationTracker(tmp_path / "job_applications.xlsx")
    job = {
        "company": "Example Co.",
        "job_title": "Frontend Developer",
        "job_location": "Remote",
        "link": "https://www.linkedin.com/jobs/view/123",
        "job_recruiter": "",
        "pdf_path": "file:///resume.pdf",
    }

    tracker.record(job, "success")
    tracker.record(job, "skipped")

    workbook = load_workbook(tracker.workbook_path)
    sheet = workbook["Applications"]
    assert sheet.max_row == 2
    assert sheet.cell(2, 1).value == "Applied"
    assert sheet.cell(2, 4).value
    assert sheet.cell(2, 5).value == "Example Co."
    assert sheet.cell(2, 8).hyperlink.target == job["link"]
    assert sheet.freeze_panes == "A2"


def test_tracker_distinguishes_unconfirmed_submission(tmp_path):
    tracker = ApplicationTracker(tmp_path / "job_applications.xlsx")
    tracker.record(
        {
            "company": "Example Co.",
            "job_title": "Frontend Developer",
            "job_location": "Remote",
            "link": "https://www.linkedin.com/jobs/view/123",
            "job_recruiter": "",
            "pdf_path": "",
        },
        "unconfirmed",
    )

    workbook = load_workbook(tracker.workbook_path)
    assert workbook["Applications"].cell(2, 1).value == "Submission Unconfirmed"


def test_tracker_stores_profile_match_score(tmp_path):
    tracker = ApplicationTracker(tmp_path / "job_applications.xlsx")
    tracker.record(
        {
            "company": "Example Co.",
            "job_title": "Frontend Developer",
            "job_location": "Remote",
            "link": "https://www.linkedin.com/jobs/view/123",
            "job_recruiter": "",
            "pdf_path": "",
            "profile_match_score": 73,
        },
        "success",
    )

    workbook = load_workbook(tracker.workbook_path)
    assert workbook["Applications"].cell(2, 11).value == 73


def test_tracker_migrates_existing_workbook_without_losing_rows(tmp_path):
    workbook_path = tmp_path / "job_applications.xlsx"
    tracker = ApplicationTracker(workbook_path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Applications"
    sheet.append(ApplicationTracker.LEGACY_HEADERS)
    sheet.append(
        [
            "Failed",
            "",
            "",
            "",
            "Example Co.",
            "Frontend Developer",
            "Remote",
            "https://www.linkedin.com/jobs/view/123",
            "",
            "",
        ]
    )
    workbook.save(workbook_path)

    tracker._load_workbook().save(workbook_path)

    migrated = load_workbook(workbook_path)
    assert migrated["Applications"].cell(1, 11).value == "Profile match (%)"
    assert migrated["Applications"].cell(2, 5).value == "Example Co."


def test_tracker_backfills_existing_json_results_without_inventing_dates(tmp_path):
    output = tmp_path / "output"
    output.mkdir()
    job = {
        "company": "Example Co.",
        "job_title": "Frontend Developer",
        "job_location": "Remote",
        "link": "https://www.linkedin.com/jobs/view/123",
        "job_recruiter": "",
        "pdf_path": "",
    }
    (output / "data.json").write_text(json.dumps([job]), encoding="utf-8")
    (output / "success.json").write_text(json.dumps([job]), encoding="utf-8")
    tracker = ApplicationTracker(output / "job_applications.xlsx")

    tracker.sync_existing_results(output)

    workbook = load_workbook(tracker.workbook_path)
    sheet = workbook["Applications"]
    assert sheet.max_row == 2
    assert sheet.cell(2, 1).value == "Applied"
    assert sheet.cell(2, 2).value in (None, "")
    assert sheet.cell(2, 4).value in (None, "")


def test_job_manager_records_application_status_in_tracker(tmp_path):
    output = tmp_path / "output"
    resume = tmp_path / "resume.pdf"
    resume.write_bytes(b"%PDF test")
    manager = AIHawkJobManager(driver=None)
    manager.set_parameters({
        "remote": False,
        "distance": 50,
        "date": {"all time": True},
        "uploads": {"resume": str(resume)},
        "outputFileDirectory": str(output),
    })
    job = SimpleNamespace(
        pdf_path=str(tmp_path / "resume.pdf"),
        company="Example Co.",
        title="Frontend Developer",
        link="https://www.linkedin.com/jobs/view/123",
        recruiter_link="",
        location="Remote",
    )

    manager.write_to_file(job, "success")

    workbook = load_workbook(output / "job_applications.xlsx")
    sheet = workbook["Applications"]
    assert sheet.max_row == 2
    assert sheet.cell(2, 1).value == "Applied"
    assert sheet.cell(2, 5).value == "Example Co."
    assert (output / "success.json").exists()
    result = json.loads((output / "success.json").read_text(encoding="utf-8"))
    assert result[0]["pdf_path"] == resume.resolve().as_uri()
