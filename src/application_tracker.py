import json
from datetime import datetime, timezone
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException


class ApplicationTrackerError(RuntimeError):
    pass


class ApplicationTracker:
    HEADERS = (
        "Status",
        "First seen (UTC)",
        "Last updated (UTC)",
        "Applied at (UTC)",
        "Company",
        "Job title",
        "Location",
        "Job URL",
        "Recruiter URL",
        "Resume PDF",
        "Profile match (%)",
    )
    LEGACY_HEADERS = HEADERS[:-1]
    STATUS_LABELS = {
        "data": "Found",
        "success": "Applied",
        "failed": "Failed",
        "unconfirmed": "Submission Unconfirmed",
        "skipped": "Skipped",
        "skipped_due_to_applicants": "Skipped",
        "skipped_low_match": "Below Match Threshold",
    }
    SOURCE_FILES = (
        "data",
        "skipped_due_to_applicants",
        "skipped",
        "failed",
        "unconfirmed",
        "success",
    )

    def __init__(self, workbook_path: Path):
        self.workbook_path = workbook_path

    @staticmethod
    def _now() -> str:
        return datetime.now(timezone.utc).isoformat(timespec="seconds")

    def sync_existing_results(self, output_directory: Path) -> None:
        try:
            self._ensure_workbook()
            for status in self.SOURCE_FILES:
                source_path = output_directory / f"{status}.json"
                if not source_path.exists():
                    continue
                with source_path.open(encoding="utf-8") as source_file:
                    records = json.load(source_file)
                if not isinstance(records, list):
                    raise ValueError(f"Expected a list in {source_path}")
                for record in records:
                    if not isinstance(record, dict):
                        raise ValueError(f"Expected job objects in {source_path}")
                    self.record(record, status, backfill=True)
        except (OSError, ValueError, KeyError, TypeError, BadZipFile, InvalidFileException) as exc:
            raise ApplicationTrackerError(
                f"Could not initialize or sync Excel tracker at {self.workbook_path}: {exc}"
            ) from exc

    def record(self, job: dict, status: str, *, backfill: bool = False) -> None:
        try:
            workbook = self._load_workbook()
            sheet = workbook["Applications"]
            link = job.get("link", "")
            row = self._find_row(sheet, link)
            now = "" if backfill else self._now()
            label = self.STATUS_LABELS.get(status, status.replace("_", " ").title())

            if row is None:
                row = sheet.max_row + 1
                sheet.cell(row=row, column=2, value=now)
                sheet.cell(row=row, column=4, value=now if label == "Applied" else "")
            elif sheet.cell(row=row, column=1).value == "Applied" and label != "Applied":
                return

            sheet.cell(row=row, column=1, value=label)
            if not backfill:
                sheet.cell(row=row, column=3, value=now)
            if label == "Applied":
                sheet.cell(row=row, column=4, value=now)

            values = (
                job.get("company", ""),
                job.get("job_title", ""),
                job.get("job_location", ""),
                link,
                job.get("job_recruiter", ""),
                job.get("pdf_path", ""),
                job.get("profile_match_score"),
            )
            for column, value in enumerate(values, start=5):
                cell = sheet.cell(row=row, column=column, value=value)
                if column in (8, 9) and value:
                    cell.hyperlink = value
                    cell.style = "Hyperlink"

            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = f"A1:K{sheet.max_row}"
            workbook.save(self.workbook_path)
        except (OSError, ValueError, KeyError, TypeError, BadZipFile, InvalidFileException) as exc:
            raise ApplicationTrackerError(
                f"Could not update Excel tracker at {self.workbook_path}: {exc}"
            ) from exc

    def _ensure_workbook(self) -> None:
        self.workbook_path.parent.mkdir(parents=True, exist_ok=True)
        if self.workbook_path.exists():
            return
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Applications"
        sheet.append(self.HEADERS)
        for column, width in enumerate((14, 24, 24, 24, 28, 36, 32, 60, 60, 60, 18), start=1):
            sheet.column_dimensions[chr(64 + column)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:K1"
        workbook.save(self.workbook_path)

    def _load_workbook(self):
        self._ensure_workbook()
        workbook = load_workbook(self.workbook_path)
        if "Applications" not in workbook.sheetnames:
            sheet = workbook.create_sheet("Applications")
            sheet.append(self.HEADERS)
        sheet = workbook["Applications"]
        headers = tuple(cell.value for cell in sheet[1])
        if headers == self.LEGACY_HEADERS:
            sheet.cell(row=1, column=len(self.HEADERS), value=self.HEADERS[-1])
            sheet.column_dimensions["K"].width = 18
        elif headers != self.HEADERS:
            raise ValueError(
                f"Unexpected columns in {self.workbook_path} on the Applications sheet"
            )
        return workbook

    @staticmethod
    def _find_row(sheet, link: str):
        if not link:
            return None
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row=row, column=8).value == link:
                return row
        return None
